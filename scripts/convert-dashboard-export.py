#!/usr/bin/env python3
"""Converts a raw "LangTech Analytics Dashboard" Google Sheet export (.xlsx)
into a products input file for scripts/push-products.mjs (see
api/_data/products.template.js for that file's shape).

Usage:
  py scripts/convert-dashboard-export.py "<path to the quarter's .xlsx export>" <quarter>
  e.g. py scripts/convert-dashboard-export.py "LangTech Analytics Dashboard_Q3.xlsx" FY26Q3

Requires `openpyxl` (py -3 -m pip install openpyxl) -- not a project
dependency since this only ever runs by hand, by the person doing the
quarterly push, not as part of the deployed API.

What it reads:

- "Product List" sheet, header on row 2 (row 1 is blank): one row per
  product, grouped under section-header rows (e.g. a lone "Literacy" in
  column A with every other column blank -- these, and the trailing
  "List of Products/Tools for further consideration" block, have a blank
  Category column). Columns are looked up BY HEADER NAME on row 2 (see
  PRODUCT_LIST_COLUMNS below), not by fixed position -- the FY26Q4 export
  inserted two new columns ("Metrics Collection Notes", "On Google Fonts?")
  ahead of previously-fixed column indices, which silently corrupted the
  platforms/devStatus/metricsMode/productUrl/openSource fields for every
  product until this was changed to header lookup. "Product Page" cells are
  frequently hyperlinked with a different, real URL than their displayed
  text -- must read the hyperlink target, not the value. "Platform(s)" is a
  comma-separated list. "Open Source Repo?" is free text ("Yes", "No",
  "Yes + LfMerge", "No, (see link)") reduced to a bool.
- One sheet per quarter (named exactly the quarter label, e.g. "FY26Q3"),
  same product-name rows, header on row 2. Also looked up by header name (see
  METRICS_COLUMNS below) except the free-text notes column, whose header
  wording changes every quarter ("Notes/Comments about metrics current
  quarter", "...about current quarter's metrics", ...) -- found instead by
  the first header starting with "Notes". The notes column feeds the
  template's `notes` field -- a short explanation of something notable about
  that product's numbers this quarter, shown on the dashboard's "What
  changed" section; blank is common and means nothing notable to call out. A
  product with no metric at all this quarter gets `metrics.pending = true`,
  so the dashboard can show a "stats not in yet" label instead of zeroes.

Inclusion rule (matches the README's "excludes the sheet's Fonts section and
any placeholder rows with no dev status and no metrics ever recorded"):
include a Product List row only if Category is non-blank and not "Fonts",
the name isn't in EXCLUDED_PRODUCTS, AND either Development Status is filled
or the row has at least one non-null metric in this quarter's sheet. Fonts
are excluded by Category rather than by a blank-Category check (as of the
FY26Q4 export, font rows have Category "Fonts" filled in, not blank) -- they
already have their own Font Usage section, fed separately by
scripts/import-font-stats.mjs, so listing them again as products would be
redundant. EXCLUDED_PRODUCTS covers products with a Development Status filled
in (so they'd otherwise pass the rule above) but that have never once had a
metric recorded, in any quarter -- confirmed empty in both FY26Q3 and FY26Q4.
"""

import difflib
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Missing dependency. Run: py -3 -m pip install openpyxl", file=sys.stderr)
    sys.exit(1)

METADATA_HEADER_ROW = 2
METADATA_FIRST_DATA_ROW = 3
METRICS_HEADER_ROW = 2
METRICS_FIRST_DATA_ROW = 3

# Column lookups below are BY HEADER NAME (see header_map/require_columns), not fixed
# position -- a past version hardcoded column indices and silently broke when the
# FY26Q4 export inserted two new columns ahead of them.
PRODUCT_LIST_COLUMNS = {
    "name": "Product Name",
    "category": "Category",
    "platforms": "Platform(s)",
    "dev_status": "Development Status",
    "metrics_mode": "Metrics collection mode",
    "product_page": "Product Page",
    "open_source": "Open Source Repo?",
}

METRICS_COLUMNS = {
    "name": "Product Name",
    "downloads": "Downloads",
    "installs": "Installs",
    "active_projects": "Active Projects",
    "active_users": "Active Users",
    "countries": "Number of User Countries",
    "languages_impacted": "Number of Languages Impacted",
}

# These have a Development Status filled in (so the usual inclusion rule would keep
# them) but have never had a single metric recorded in any quarter sheet -- confirmed
# empty in both FY26Q3 and FY26Q4. Excluded by name rather than waiting on the
# Development Status column, since nothing else distinguishes them from an active,
# just-not-yet-reporting product.
EXCLUDED_PRODUCTS = {
    "Reading App Builder",
    "Dictionary App Builder",
    "Scripture App Builder",
    "Keyboard App Builder",
}


def header_map(ws, row):
    headers = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=row, column=col).value
        if isinstance(value, str) and value.strip():
            headers.setdefault(value.strip(), col)
    return headers


def require_columns(headers, labels, sheet_name, row):
    missing = [label for label in labels if label not in headers]
    if missing:
        sys.exit(
            f"{sheet_name}: expected column(s) {missing} not found in header row {row}. "
            f"Found headers: {sorted(headers)}. The sheet layout has changed -- "
            f"update convert-dashboard-export.py's column mapping."
        )


def closest_name(name, candidates):
    # Plain difflib similarity ranks by overall string shape, so for short names it can
    # favor an unrelated same-length candidate ("Scripture App Builder") over the actual
    # near-duplicate ("Scriptoria", inside "Scriptoria App Service") -- prefer a
    # substring match either direction before falling back to difflib.
    name_lower = name.lower()
    substring_matches = [c for c in candidates if c.lower() in name_lower or name_lower in c.lower()]
    if substring_matches:
        return max(substring_matches, key=len)
    close = difflib.get_close_matches(name, candidates, n=1, cutoff=0.6)
    return close[0] if close else None


def notes_column(headers):
    # The notes header's exact wording changes every quarter -- match by prefix
    # instead of an exact label. None (rather than exiting) if no such column exists;
    # notes are optional commentary, not something worth failing the whole import over.
    for label, col in headers.items():
        if label.lower().startswith("notes"):
            return col
    return None


def clean_number(value):
    # Metric cells are sometimes a placeholder string ("N/A", "n/a", "-", ...)
    # instead of being left blank -- treat anything that isn't a real number
    # as no data, rather than letting a literal "n/a" string reach the JSON
    # (that's what showed up as NaN on the front end for Scripture Forge and
    # Serval's active_users/countries).
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            number = float(text)
        except ValueError:
            return None
        return int(number) if number.is_integer() else number
    return None


def clean_text(value):
    # Notes/Comments is free text -- blank is the common case (nothing notable to
    # call out), so normalize both None and whitespace-only to None rather than "".
    if not isinstance(value, str):
        return None
    text = value.strip()
    return text or None


def read_metrics_sheet(ws):
    headers = header_map(ws, METRICS_HEADER_ROW)
    require_columns(headers, METRICS_COLUMNS.values(), ws.title, METRICS_HEADER_ROW)
    cols = {key: headers[label] for key, label in METRICS_COLUMNS.items()}
    notes_col = notes_column(headers)

    by_name = {}
    for row in range(METRICS_FIRST_DATA_ROW, ws.max_row + 1):
        name = ws.cell(row=row, column=cols["name"]).value
        if not name:
            continue
        metrics = {
            "downloads": clean_number(ws.cell(row=row, column=cols["downloads"]).value),
            "installs": clean_number(ws.cell(row=row, column=cols["installs"]).value),
            "active_projects": clean_number(ws.cell(row=row, column=cols["active_projects"]).value),
            "active_users": clean_number(ws.cell(row=row, column=cols["active_users"]).value),
            "countries": clean_number(ws.cell(row=row, column=cols["countries"]).value),
            "languages_impacted": clean_number(ws.cell(row=row, column=cols["languages_impacted"]).value),
            "notes": clean_text(ws.cell(row=row, column=notes_col).value) if notes_col else None,
        }
        metrics["pending"] = not has_any_metric(metrics)
        by_name[name] = metrics
    return by_name


def product_url(cell):
    if cell.hyperlink is not None and cell.hyperlink.target:
        return cell.hyperlink.target
    value = cell.value
    if isinstance(value, str) and value.lower().startswith("http"):
        return value
    return None


def is_open_source(value):
    if not isinstance(value, str):
        return False
    return value.strip().lower().startswith("yes")


def has_any_metric(metrics):
    # "notes" is commentary, not a usage metric -- a stray comment on an otherwise
    # all-blank row shouldn't by itself count as "this product reported data".
    return any(v is not None for k, v in metrics.items() if k not in ("notes", "pending"))


def empty_metrics():
    return {
        "downloads": None, "installs": None, "active_projects": None,
        "active_users": None, "countries": None, "languages_impacted": None,
        "notes": None, "pending": True,
    }


def main():
    if len(sys.argv) != 3:
        print(f"Usage: py {sys.argv[0]} <path to .xlsx export> <quarter, e.g. FY26Q3>", file=sys.stderr)
        sys.exit(1)

    xlsx_path = Path(sys.argv[1])
    quarter = sys.argv[2]

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if quarter not in wb.sheetnames:
        print(f"No sheet named '{quarter}' in {xlsx_path.name}. Sheets: {wb.sheetnames}", file=sys.stderr)
        sys.exit(1)

    metrics_by_name = read_metrics_sheet(wb[quarter])

    product_list = wb["Product List"]
    pl_headers = header_map(product_list, METADATA_HEADER_ROW)
    require_columns(pl_headers, PRODUCT_LIST_COLUMNS.values(), "Product List", METADATA_HEADER_ROW)
    pl = {key: pl_headers[label] for key, label in PRODUCT_LIST_COLUMNS.items()}

    products = []
    skipped = []
    excluded = []
    unmatched_names = []
    for row in range(METADATA_FIRST_DATA_ROW, product_list.max_row + 1):
        name = product_list.cell(row=row, column=pl["name"]).value
        if not name:
            continue
        category = product_list.cell(row=row, column=pl["category"]).value
        if not category:
            continue  # section headers, "further consideration" block
        if category == "Fonts":
            continue  # fonts have their own Font Usage section, fed by import-font-stats.mjs
        if name in EXCLUDED_PRODUCTS:
            excluded.append(name)
            continue

        dev_status = product_list.cell(row=row, column=pl["dev_status"]).value
        metrics = metrics_by_name.get(name, empty_metrics())
        if not dev_status and not has_any_metric(metrics):
            skipped.append(name)
            continue
        if name not in metrics_by_name:
            # A Product List name with no exact match in the quarter sheet silently reads
            # as "no metrics this quarter" even if the row is there under a slightly
            # different spelling (e.g. "Scriptoria App Service" vs "Scriptoria") -- warn
            # instead of letting that data vanish unnoticed.
            unmatched_names.append(name)

        platforms_raw = product_list.cell(row=row, column=pl["platforms"]).value or ""
        platforms = [p.strip() for p in platforms_raw.split(",") if p.strip()]
        metrics_mode = product_list.cell(row=row, column=pl["metrics_mode"]).value

        products.append({
            "name": name,
            "category": category,
            "platforms": platforms,
            "devStatus": dev_status,
            "metricsMode": metrics_mode,
            "productUrl": product_url(product_list.cell(row=row, column=pl["product_page"])),
            "openSource": is_open_source(product_list.cell(row=row, column=pl["open_source"]).value),
            "metrics": metrics,
        })

    from datetime import datetime, timezone
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")

    output = {
        "quarter": quarter,
        "generatedAt": generated_at,
        "products": products,
    }

    repo_root = Path(__file__).resolve().parent.parent
    out_dir = repo_root / "input" / "products"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{quarter}.js"
    body = json.dumps(output, indent=2, ensure_ascii=False)
    out_path.write_text(
        f"// Auto-generated by scripts/convert-dashboard-export.py from\n"
        f"// {xlsx_path.name}. Do not hand-edit -- re-run the converter instead.\n"
        f"export default {body};\n",
        encoding="utf-8",
    )

    print(f"Wrote {out_path} ({len(products)} product(s)).")
    if skipped:
        print(f"Skipped {len(skipped)} placeholder row(s) (no dev status, no metrics): {', '.join(skipped)}")
    if excluded:
        print(f"Excluded {len(excluded)} product(s) with no metrics ever recorded: {', '.join(excluded)}")
    if unmatched_names:
        print(
            f"WARNING: {len(unmatched_names)} product(s) from 'Product List' have no exact-name "
            f"match in the '{quarter}' sheet, so they'll show as not-yet-reported even if that "
            f"sheet has their numbers under a different spelling:",
            file=sys.stderr,
        )
        for name in unmatched_names:
            match = closest_name(name, list(metrics_by_name.keys()))
            suggestion = f" (closest match in '{quarter}': '{match}')" if match else ""
            print(f"  - {name!r}{suggestion}", file=sys.stderr)


if __name__ == "__main__":
    main()
